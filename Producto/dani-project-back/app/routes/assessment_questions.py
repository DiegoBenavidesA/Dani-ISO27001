import asyncio
import io
import os
import tempfile
import json
import re
import unicodedata
from typing import List, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies.database import get_db
from app.dependencies.auth import get_current_user, get_current_org, RequireRole, ELEVATED_WRITE
from app.models.assessment_question import AssessmentQuestion
from app.models.assessment_answer import AssessmentAnswer
from app.models.control_status import ControlStatus


def _codigo_a_control_id(codigo: str) -> str:
    """Convierte el código de una pregunta al control_id ISO.

    Las preguntas usan "A.5.1" (Anexo A) y los controles ISO usan "5.1".
    """
    c = (codigo or "").strip().upper()
    if c.startswith("A."):
        c = c[2:]
    elif c.startswith("A") and len(c) > 1 and c[1].isdigit():
        c = c[1:]
    return c.strip(". ")
from app.services.ai_service import AIService

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

router = APIRouter(prefix="/api/assessment-questions", tags=["Assessment Questions"])
ai_service = AIService()

class QuestionResponse(BaseModel):
    id: str
    codigo: str
    categoria: str
    nombre: str
    pregunta: str
    evidencia_esperada: Optional[str] = None
    orden: int
    class Config: from_attributes = True

class QuestionCreate(BaseModel):
    codigo: str
    categoria: str
    nombre: str
    pregunta: str
    evidencia_esperada: Optional[str] = None
    orden: int = 0

class QuestionUpdate(BaseModel):
    codigo: Optional[str] = None
    categoria: Optional[str] = None
    nombre: Optional[str] = None
    pregunta: Optional[str] = None
    evidencia_esperada: Optional[str] = None
    orden: Optional[int] = None

class AnswerResponse(BaseModel):
    id: str
    question_id: str
    organization_id: str
    veredicto: str
    confianza: Optional[float] = None
    justificacion: Optional[str] = None
    class Config: from_attributes = True

@router.get("/", response_model=List[QuestionResponse])
async def get_questions(categoria: Optional[str] = None, current_user: dict = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    stmt = select(AssessmentQuestion)
    if categoria: stmt = stmt.where(AssessmentQuestion.categoria == categoria)
    stmt = stmt.order_by(AssessmentQuestion.orden)
    result = await db.execute(stmt)
    return result.scalars().all()

@router.post("/", response_model=QuestionResponse)
async def create_question(data: QuestionCreate, current_user: dict = Depends(RequireRole(["admin"])), db: AsyncSession = Depends(get_db)):
    new_q = AssessmentQuestion(**data.model_dump(exclude_unset=True))
    db.add(new_q)
    await db.commit()
    await db.refresh(new_q)
    return new_q

@router.put("/{q_id}", response_model=QuestionResponse)
async def update_question(q_id: str, data: QuestionUpdate, current_user: dict = Depends(RequireRole(["admin"])), db: AsyncSession = Depends(get_db)):
    q = (await db.execute(select(AssessmentQuestion).where(AssessmentQuestion.id == q_id))).scalar_one_or_none()
    if not q: raise HTTPException(status_code=404, detail="Pregunta no encontrada")
    for key, value in data.model_dump(exclude_unset=True).items(): setattr(q, key, value)
    await db.commit()
    await db.refresh(q)
    return q

@router.delete("/{q_id}")
async def delete_question(q_id: str, current_user: dict = Depends(RequireRole(["admin"])), db: AsyncSession = Depends(get_db)):
    q = (await db.execute(select(AssessmentQuestion).where(AssessmentQuestion.id == q_id))).scalar_one_or_none()
    if not q: raise HTTPException(status_code=404, detail="Pregunta no encontrada")
    await db.delete(q)
    await db.commit()
    return {"message": "Eliminada"}


# =========================================================
# CARGA MASIVA DE PREGUNTAS DESDE EXCEL (.xlsx)
# Mapea columnas de forma flexible: acepta otros nombres y columnas de más.
# Solo 'pregunta' es obligatoria; el resto usa valores por defecto.
# =========================================================
_COLUMN_SYNONYMS = {
    "pregunta": ["pregunta", "question", "cuestion", "enunciado", "consulta", "textopregunta", "detalle", "item"],
    "categoria": ["categoria", "category", "dominio", "domain", "seccion", "section", "grupo", "area", "marco", "norma", "anexo", "tipo"],
    "codigo": ["codigo", "code", "controlid", "idcontrol", "codigocontrol", "clausula", "referencia", "ref", "numerocontrol", "numero", "nro", "id", "control"],
    "nombre": ["nombre", "name", "titulo", "title", "nombrecontrol", "nombredelcontrol", "nombrecorto", "descripcioncorta"],
    "evidencia_esperada": ["evidenciaesperada", "evidencia", "evidence", "evidenciarequerida", "expectedevidence", "pruebaesperada"],
    "orden": ["orden", "order", "posicion", "position", "secuencia", "indice", "nroorden"],
}


def _normalizar(texto) -> str:
    if texto is None:
        return ""
    s = str(texto).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", s)


def _mapear_columnas(headers: list) -> dict:
    norm = [_normalizar(h) for h in headers]
    usados = set()
    mapeo = {}
    for campo, sinonimos in _COLUMN_SYNONYMS.items():
        for syn in sinonimos:
            for i, h in enumerate(norm):
                if i in usados or not h:
                    continue
                if h == syn:
                    mapeo[campo] = i
                    usados.add(i)
                    break
            if campo in mapeo:
                break
    for campo, sinonimos in _COLUMN_SYNONYMS.items():
        if campo in mapeo:
            continue
        for syn in sinonimos:
            for i, h in enumerate(norm):
                if i in usados or not h:
                    continue
                if syn in h:
                    mapeo[campo] = i
                    usados.add(i)
                    break
            if campo in mapeo:
                break
    return mapeo


@router.post("/bulk-import")
async def bulk_import_questions(
    file: UploadFile = File(...),
    current_user: dict = Depends(RequireRole(["admin"])),
    db: AsyncSession = Depends(get_db),
):
    """Carga masiva de preguntas desde un Excel (.xlsx). Tolerante a nombres de columna distintos."""
    nombre_archivo = (file.filename or "").lower()
    if not nombre_archivo.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel (.xlsx).")

    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="Falta la dependencia 'openpyxl' en el servidor.")

    contenido = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el Excel. ¿Está dañado o no es un .xlsx válido?")

    ws = wb.active
    filas = ws.iter_rows(values_only=True)

    headers = None
    for fila in filas:
        if fila and any(c is not None and str(c).strip() != "" for c in fila):
            headers = list(fila)
            break
    if not headers:
        raise HTTPException(status_code=400, detail="El Excel está vacío.")

    mapeo = _mapear_columnas(headers)
    if "pregunta" not in mapeo:
        raise HTTPException(
            status_code=400,
            detail="No se encontró la columna de preguntas. Agrega una columna llamada 'pregunta' (o similar).",
        )

    def val(fila, campo):
        idx = mapeo.get(campo)
        if idx is None or idx >= len(fila):
            return None
        v = fila[idx]
        if v is None:
            return None
        v = str(v).strip()
        return v or None

    existentes = set()
    res_prev = await db.execute(select(AssessmentQuestion.codigo, AssessmentQuestion.pregunta))
    for cod, preg in res_prev.all():
        existentes.add((_normalizar(cod), _normalizar(preg)))

    creadas = 0
    omitidas = 0
    detalles_omitidas = []
    nuevas_en_lote = set()
    numero_fila = 1

    for fila in filas:
        numero_fila += 1
        if not fila or all(c is None or str(c).strip() == "" for c in fila):
            continue
        pregunta = val(fila, "pregunta")
        if not pregunta:
            omitidas += 1
            detalles_omitidas.append({"fila": numero_fila, "motivo": "sin pregunta"})
            continue
        codigo = (val(fila, "codigo") or "S/C")[:50]
        categoria = (val(fila, "categoria") or "General")[:100]
        nombre = (val(fila, "nombre") or pregunta[:60])[:255]
        evidencia = val(fila, "evidencia_esperada")
        orden_raw = val(fila, "orden")
        try:
            orden = int(float(orden_raw)) if orden_raw else 0
        except (ValueError, TypeError):
            orden = 0

        clave = (_normalizar(codigo), _normalizar(pregunta))
        if clave in existentes or clave in nuevas_en_lote:
            omitidas += 1
            detalles_omitidas.append({"fila": numero_fila, "motivo": "duplicada"})
            continue

        db.add(AssessmentQuestion(
            codigo=codigo, categoria=categoria, nombre=nombre,
            pregunta=pregunta, evidencia_esperada=evidencia, orden=orden,
        ))
        nuevas_en_lote.add(clave)
        creadas += 1

    if creadas:
        await db.commit()

    columnas_detectadas = {campo: headers[idx] for campo, idx in mapeo.items() if idx < len(headers)}
    return {
        "creadas": creadas,
        "omitidas": omitidas,
        "total_filas_datos": creadas + omitidas,
        "columnas_detectadas": columnas_detectadas,
        "detalles_omitidas": detalles_omitidas[:50],
    }


@router.get("/bulk-template")
async def bulk_template(current_user: dict = Depends(RequireRole(["admin"]))):
    """Descarga una plantilla .xlsx con las columnas esperadas y 2 ejemplos."""
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="Falta la dependencia 'openpyxl' en el servidor.")
    from fastapi.responses import StreamingResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Preguntas"
    ws.append(["codigo", "categoria", "nombre", "pregunta", "evidencia_esperada", "orden"])
    ws.append([
        "A.5.1", "Organizacional", "Políticas de seguridad",
        "¿La organización tiene políticas de seguridad aprobadas?",
        "Documento de política firmado por la dirección", 1,
    ])
    ws.append([
        "O1", "Privacidad", "Registro de tratamientos (RoPA)",
        "¿Existe un registro de actividades de tratamiento de datos personales?",
        "Inventario RoPA actualizado", 2,
    ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_preguntas.xlsx"},
    )

@router.get("/answers", response_model=List[AnswerResponse])
async def get_answers(org_id: str = Depends(get_current_org), db: AsyncSession = Depends(get_db)):
    stmt = select(AssessmentAnswer).where(AssessmentAnswer.organization_id == org_id).order_by(AssessmentAnswer.created_at)
    return (await db.execute(stmt)).scalars().all()

@router.get("/answers/{question_id}", response_model=AnswerResponse)
async def get_answer_by_question(question_id: str, org_id: str = Depends(get_current_org), db: AsyncSession = Depends(get_db)):
    answer = (await db.execute(select(AssessmentAnswer).where(AssessmentAnswer.question_id == question_id, AssessmentAnswer.organization_id == org_id))).scalar_one_or_none()
    if not answer: raise HTTPException(status_code=404, detail="No encontrada")
    return answer

# Función para extraer JSON de una cadena de texto (incluso si la IA agrega markdown o comentarios)
def extract_json_from_text(text: str):
    try:
        # Intenta parsear directamente
        return json.loads(text)
    except json.JSONDecodeError:
        # Si falla, busca un bloque JSON dentro del texto
        try:
            match = re.search(r'\[\s*\{.*?\}\s*\]', text, re.DOTALL | re.MULTILINE)
            if match:
                return json.loads(match.group(0))
            
            # Intenta buscar un objeto JSON simple si no es un arreglo
            match_obj = re.search(r'\{\s*".*?:.*?\}', text, re.DOTALL | re.MULTILINE)
            if match_obj:
                 return [json.loads(match_obj.group(0))]

            return []
        except:
            print(f"No se pudo extraer JSON de la respuesta de la IA: {text}")
            return []

# EVALUACIÓN EXHAUSTIVA DE IA Y RoPA AUTOMÁTICA
@router.post("/evaluate")
async def evaluate_with_ai(
    files: List[UploadFile] = File(...),
    question_ids: str = Form(""),
    current_user: dict = Depends(RequireRole(ELEVATED_WRITE)),
    org_id: str = Depends(get_current_org),
    db: AsyncSession = Depends(get_db)
):
    contexto = ""
    for f in files:
        data = await f.read()
        texto = ""
        filename = f.filename.lower() if f.filename else ""
        
        # Extracción robusta de PDF
        if filename.endswith('.pdf') and PdfReader:
            try:
                import io
                pdf = PdfReader(io.BytesIO(data))
                for page in pdf.pages:
                    extracted = page.extract_text()
                    if extracted: texto += extracted + "\n"
            except Exception as e:
                print(f"Error leyendo PDF: {e}")
        else:
            try: texto = data.decode('utf-8')
            except: texto = data.decode('latin-1', errors='ignore')
            
        if texto.strip():
            contexto += f"\n\n### Documento: {f.filename}\n{texto}"

    if not contexto.strip(): raise HTTPException(status_code=400, detail="No se pudo extraer texto. Asegúrate de que los PDFs no sean solo imágenes.")

    stmt = select(AssessmentQuestion)
    ids = [x.strip() for x in question_ids.split(",") if x.strip()]
    if ids: stmt = stmt.where(AssessmentQuestion.id.in_(ids))
    preguntas = (await db.execute(stmt.order_by(AssessmentQuestion.orden))).scalars().all()

    resultados = []
    if preguntas:
        TAMANO_LOTE = 10
        lotes = [preguntas[i:i + TAMANO_LOTE] for i in range(0, len(preguntas), TAMANO_LOTE)]
        sem = asyncio.Semaphore(3)

        async def evaluar_lote(qs):
            async with sem:
                prompt = f"""Eres un Auditor ISO 27001 y Ley 21.719. Analiza este documento:
                {contexto[:8000]}
                
                Responde a estas preguntas basadas SOLAMENTE en el documento:
                {json.dumps([{"id": q.id, "pregunta": q.pregunta, "codigo": q.codigo} for q in qs])}
                
                Instrucciones:
                1. Si el documento muestra explícitamente que cumplen la pregunta, "veredicto" es "cumple".
                2. Si el documento muestra cumplimiento parcial o falta información clave, "veredicto" es "parcial".
                3. Si el documento no menciona el tema o no hay evidencia, "veredicto" es "sin_evidencia".
                4. "justificacion" debe ser una breve razón de tu veredicto citando el documento.
                
                MUY IMPORTANTE: Tu respuesta DEBE ser un arreglo JSON válido y NADA MÁS. No incluyas bloques de código markdown (```json), ni texto introductorio, ni comentarios.
                
                Ejemplo de formato esperado:
                [
                  {{ "id": "id_pregunta_1", "veredicto": "cumple", "confianza": 0.9, "justificacion": "El documento menciona..." }},
                  {{ "id": "id_pregunta_2", "veredicto": "sin_evidencia", "confianza": 0.1, "justificacion": "No se menciona en el texto proporcionado." }}
                ]
                """
                resp_text = await ai_service.generate_document(prompt)
                
                parsed_json = extract_json_from_text(resp_text)
                if parsed_json:
                    return parsed_json
                else:
                    return [{"id": q.id, "veredicto": "sin_evidencia", "confianza": 0.0, "justificacion": "Error: La IA no devolvió un JSON válido."} for q in qs]

        lotes_res = await asyncio.gather(*[evaluar_lote(qs) for qs in lotes])
        resultados = [item for sub in lotes_res for item in sub if type(item) is dict]

        q_ids_res = [str(item.get("id")) for item in resultados if item.get("id")]
        respuestas_existentes = {}
        if q_ids_res:
            exist_stmt = select(AssessmentAnswer).where(AssessmentAnswer.organization_id == org_id, AssessmentAnswer.question_id.in_(q_ids_res))
            respuestas_existentes = {a.question_id: a for a in (await db.execute(exist_stmt)).scalars().all()}

        for item in resultados:
            qid = str(item.get("id") or "").strip()
            if not qid: continue
            veredicto = item.get("veredicto", "sin_evidencia")
            if veredicto not in {"cumple", "parcial", "no_cumple", "sin_evidencia"}: veredicto = "sin_evidencia"
            try: confianza = float(item.get("confianza", 0.0))
            except: confianza = 0.0
            
            existing = respuestas_existentes.get(qid)
            if existing:
                existing.veredicto = veredicto
                existing.confianza = confianza
                existing.justificacion = item.get("justificacion")
            else:
                db.add(AssessmentAnswer(question_id=qid, organization_id=org_id, veredicto=veredicto, confianza=confianza, justificacion=item.get("justificacion")))
        await db.commit()

        # --- Puente hacia el SOA: actualizar el estado de los controles ISO ---
        # El veredicto de la IA se refleja en ControlStatus para que el SOA y
        # Resultados cambien. Mapeo por código: pregunta "A.5.1" -> control "5.1".
        q_by_id = {q.id: q for q in preguntas}
        estado_por_veredicto = {
            "cumple": ("Implementado", 100),
            "parcial": ("Planificado", 50),
            "no_cumple": ("No Implementado", 0),
            "sin_evidencia": ("No Implementado", 0),
        }
        for item in resultados:
            qid = str(item.get("id") or "").strip()
            q = q_by_id.get(qid)
            if not q:
                continue
            control_id = _codigo_a_control_id(q.codigo)
            veredicto = item.get("veredicto", "sin_evidencia")
            nuevo_estado, score = estado_por_veredicto.get(veredicto, ("No Implementado", 0))
            # Solo tocamos controles que existan para esta empresa.
            cs_res = await db.execute(
                select(ControlStatus).where(
                    ControlStatus.control_id == control_id,
                    ControlStatus.organization_id == org_id,
                )
            )
            cs = cs_res.scalar_one_or_none()
            if not cs:
                continue
            # No degradamos un control ya Implementado por un "sin_evidencia".
            if veredicto in ("sin_evidencia", "no_cumple") and cs.status == "Implementado":
                continue
            cs.applies = True
            cs.status = nuevo_estado
            cs.score = score
            if item.get("justificacion"):
                cs.justification = item.get("justificacion")
        await db.commit()

    # EXTRAER TRATAMIENTOS RoPA AUTOMÁTICAMENTE
    extracted_treatments = []
    try:
        ropa_prompt = f"""Analiza detalladamente este documento y extrae TODAS las Actividades de Tratamiento de Datos Personales (para un Registro RoPA).
        Documento: 
        {contexto[:8000]}
        
        Busca información como: gestión de clientes, marketing, recursos humanos, videovigilancia, etc.
        Para cada tratamiento encontrado, extrae la siguiente información:
        - "nombre": Nombre descriptivo de la actividad (ej. "Gestión de clientes y facturación").
        - "finalidad": Para qué se usan los datos.
        - "base_licitud": La base legal (ej. "ejecución de un contrato", "consentimiento").
        - "categorias_datos": Qué datos se recopilan (ej. "nombre, RUT, dirección").
        - "plazo_conservacion": Cuánto tiempo se guardan.
        - "destinatarios": A quién se comparten los datos (opcional, si se menciona).
        
        Instrucciones MUY IMPORTANTES:
        1. Devuelve ÚNICAMENTE un arreglo JSON válido y NADA MÁS.
        2. No incluyas bloques de código markdown (```json), ni explicaciones.
        3. Si no encuentras NINGÚN tratamiento, devuelve un arreglo vacío: []
        
        Ejemplo de formato esperado:
        [
          {{
            "nombre": "Envío de comunicaciones de marketing",
            "finalidad": "enviar newsletters y promociones",
            "base_licitud": "consentimiento del titular",
            "categorias_datos": "nombre y correo electrónico",
            "plazo_conservacion": "hasta que revoque el consentimiento",
            "destinatarios": "plataforma de email marketing"
          }}
        ]
        """
        
        ropa_text = await ai_service.generate_document(ropa_prompt)
        
        extracted_treatments = extract_json_from_text(ropa_text)
        print(f"Tratamientos extraídos: {len(extracted_treatments)}")

    except Exception as e:
        print(f"Error en la extracción de RoPA: {e}")

    return { "total": len(resultados), "results": resultados, "extracted_treatments": extracted_treatments }